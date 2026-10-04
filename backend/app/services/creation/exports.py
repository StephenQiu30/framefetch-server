"""Portable revision-bound files; no network fetching or platform account writes."""

from __future__ import annotations

import csv
import hashlib
import html
import json
import re
from collections.abc import Mapping, Sequence
from dataclasses import dataclass, field
from html.parser import HTMLParser
from io import BytesIO, StringIO
from pathlib import PurePosixPath
from typing import Literal
from urllib.parse import unquote, urlsplit
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from markdown_it import MarkdownIt
from PIL import Image

ExportFormat = Literal["md", "docx", "html", "zip"]
_ALLOWED = frozenset(
    "p h1 h2 h3 h4 h5 h6 strong em b i s blockquote pre code ul ol li "
    "table thead tbody tr th td a img hr br span".split()
)
_DROP = frozenset("script style iframe object embed svg math template".split())
_VOID = frozenset({"img", "hr", "br"})
_STYLE = {
    "p": "margin:0 0 16px;line-height:1.8;overflow-wrap:anywhere",
    "blockquote": "margin:16px 0;padding-left:16px;border-left:3px solid #999",
    "pre": "white-space:pre-wrap;overflow-wrap:anywhere",
    "img": "max-width:100%;height:auto",
    "table": "border-collapse:collapse;max-width:100%",
}


@dataclass(frozen=True, slots=True)
class ExportAsset:
    name: str
    data: bytes
    media_type: str
    rights: str


@dataclass(frozen=True, slots=True)
class ExportDocument:
    title: str
    body: str
    revision_id: str
    references: Sequence[Mapping[str, object]] = ()
    assets: Sequence[ExportAsset] = ()


@dataclass(frozen=True, slots=True)
class ExportArtifact:
    filename: str
    media_type: str
    data: bytes
    metadata: dict[str, object] = field(default_factory=dict)


def _safe_name(name: str) -> str:
    decoded = unquote(name)
    path = PurePosixPath(decoded)
    if (
        not name
        or len(name) > 160
        or name != decoded
        or path.name != name
        or name in (".", "..")
        or any(char in name for char in "/\\:\x00\r\n")
        or any(ord(char) < 32 for char in name)
    ):
        raise ValueError("素材文件名必须为安全的单个文件名")
    return name


def _check_document(document: ExportDocument) -> None:
    if not document.title.strip() or not document.revision_id.strip():
        raise ValueError("导出需要标题和确认版本")
    if len(document.title) > 500 or len(document.body) > 30_000:
        raise ValueError("标题或正文超过导出范围")
    if "\x00" in document.title or "\x00" in document.body:
        raise ValueError("正文不能含空字节")
    if len(document.assets) > 40 or len(document.references) > 200:
        raise ValueError("素材或引用超过导出范围")
    names: set[str] = set()
    total_size = 0
    for asset in document.assets:
        name = _safe_name(asset.name)
        if name.casefold() in names:
            raise ValueError("素材文件名重复")
        names.add(name.casefold())
        if not asset.rights.strip() or asset.rights.strip().lower() in {
            "unknown",
            "unconfirmed",
            "未知",
            "待确认",
            "许可不明",
        }:
            raise ValueError("素材权利未确认，不能打包分发")
        expected = {
            "image/png": "PNG",
            "image/jpeg": "JPEG",
            "image/webp": "WEBP",
        }.get(asset.media_type)
        if expected is None:
            raise ValueError("导出素材仅支持实际 PNG/JPEG/WebP")
        total_size += len(asset.data)
        if not asset.data or len(asset.data) > 10 * 1024 * 1024:
            raise ValueError("素材为空或超过 10 MiB")
        try:
            with Image.open(BytesIO(asset.data)) as image:
                if (
                    image.format != expected
                    or image.width * image.height > 20_000_000
                    or getattr(image, "n_frames", 1) != 1
                ):
                    raise ValueError("素材编码或像素范围无效")
                image.verify()
        except (OSError, Image.DecompressionBombError) as exc:
            raise ValueError("素材不是可读取的图片") from exc
    if total_size > 64 * 1024 * 1024:
        raise ValueError("素材总大小超过 64 MiB")


def _safe_link(value: str) -> bool:
    if any(char.isspace() or ord(char) < 32 for char in value):
        return False
    try:
        parsed = urlsplit(html.unescape(value))
    except ValueError:
        return False
    return (
        parsed.scheme.lower() in ("http", "https", "mailto")
        and not parsed.username
        and not parsed.password
        and (bool(parsed.netloc) or parsed.scheme.lower() == "mailto")
    )


class _SafeHTML(HTMLParser):
    def __init__(self, allowed_assets: frozenset[str]) -> None:
        super().__init__(convert_charrefs=True)
        self.allowed_assets = allowed_assets
        self.parts: list[str] = []
        self.blocked: list[str] = []
        self.open_tags: list[str] = []
        self.missing_asset_count = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in _DROP:
            self.blocked.append(tag)
            return
        if self.blocked or tag not in _ALLOWED:
            return
        values = dict(attrs)
        safe_attrs: dict[str, str] = {}
        if tag == "a" and (href := values.get("href")) and _safe_link(href):
            safe_attrs["href"] = href
            safe_attrs["rel"] = "noopener noreferrer"
        if tag == "img":
            source = values.get("src", "")
            if source not in self.allowed_assets:
                self.missing_asset_count += 1
                self.parts.append("<span>【图片未包含在已确认素材包】</span>")
                return
            safe_attrs["src"] = source
            safe_attrs["alt"] = values.get("alt") or ""
        if tag in _STYLE:
            safe_attrs["style"] = _STYLE[tag]
        attributes = "".join(
            f' {name}="{html.escape(value, quote=True)}"'
            for name, value in safe_attrs.items()
        )
        self.parts.append(f"<{tag}{attributes}>")
        if tag not in _VOID:
            self.open_tags.append(tag)

    def handle_startendtag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        self.handle_starttag(tag, attrs)
        if tag not in _VOID:
            self.handle_endtag(tag)

    def handle_endtag(self, tag: str) -> None:
        if self.blocked:
            if tag == self.blocked[-1]:
                self.blocked.pop()
            return
        if tag in self.open_tags:
            while self.open_tags:
                current = self.open_tags.pop()
                self.parts.append(f"</{current}>")
                if current == tag:
                    break

    def handle_data(self, data: str) -> None:
        if not self.blocked:
            self.parts.append(html.escape(data))

    def finish(self) -> str:
        while self.open_tags:
            self.parts.append(f"</{self.open_tags.pop()}>")
        return "".join(self.parts)


def sanitize_html(content: str, *, allowed_assets: frozenset[str] = frozenset()) -> str:
    if len(content) > 1_000_000:
        raise ValueError("HTML 超过安全预览范围")
    parser = _SafeHTML(allowed_assets)
    parser.feed(content)
    parser.close()
    return parser.finish()


def _artifact(
    document: ExportDocument, extension: str, media: str, data: bytes
) -> ExportArtifact:
    return ExportArtifact(
        f"content.{extension}",
        media,
        data,
        {
            "revision_id": document.revision_id,
            "sha256": hashlib.sha256(data).hexdigest(),
        },
    )


def export_markdown(document: ExportDocument) -> ExportArtifact:
    _check_document(document)
    title = document.title.replace("\n", " ").replace("\r", " ")
    content = f"# {title}\n\n{document.body}\n"
    return _artifact(
        document, "md", "text/markdown; charset=utf-8", content.encode("utf-8")
    )


def export_docx(document: ExportDocument) -> ExportArtifact:
    _check_document(document)
    result = Document()
    result.core_properties.identifier = document.revision_id
    result.core_properties.title = document.title
    result.add_heading(document.title, 0)
    # Preserve all source text: markup stays editable instead of lossy HTML import.
    for paragraph in document.body.split("\n\n"):
        result.add_paragraph(paragraph)
    stream = BytesIO()
    result.save(stream)
    return _artifact(
        document,
        "docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        stream.getvalue(),
    )


def export_html(document: ExportDocument) -> ExportArtifact:
    _check_document(document)
    renderer = MarkdownIt("commonmark", {"html": False, "breaks": True})
    allowed_assets = frozenset(f"assets/{asset.name}" for asset in document.assets)
    parser = _SafeHTML(allowed_assets)
    parser.feed(renderer.render(document.body))
    parser.close()
    body = parser.finish()
    title = html.escape(document.title)
    content = (
        '<!doctype html><html lang="zh-CN"><head><meta charset="utf-8">'
        '<meta name="viewport" content="width=device-width, initial-scale=1">'
        '<meta http-equiv="Content-Security-Policy" '
        'content="default-src &#39;none&#39;; img-src &#39;self&#39;; '
        'style-src &#39;unsafe-inline&#39;; base-uri &#39;none&#39;">'
        f"<title>{title}</title></head><body><h1>{title}</h1>{body}</body></html>"
    )
    result = _artifact(
        document, "html", "text/html; charset=utf-8", content.encode("utf-8")
    )
    result.metadata.update(
        {
            "missing_asset_count": parser.missing_asset_count,
            "resource_status": "needs_assets"
            if parser.missing_asset_count
            else "complete",
            "platform_handoff_verified": False,
        }
    )
    return result


def export_zip(document: ExportDocument) -> ExportArtifact:
    _check_document(document)
    preview = export_html(document)
    if preview.metadata["missing_asset_count"]:
        raise ValueError("正文有缺失或未经确认的图片；请补充素材或移除引用后再打包")
    files = {
        item.filename: item.data
        for item in (export_markdown(document), preview, export_docx(document))
    }
    for asset in document.assets:
        files[f"assets/{asset.name}"] = asset.data
    manifest = {
        "revision_id": document.revision_id,
        "title": document.title,
        "references": list(document.references),
        "files": [
            {
                "path": name,
                "sha256": hashlib.sha256(data).hexdigest(),
                "size": len(data),
            }
            for name, data in files.items()
        ],
        "assets": [
            {
                "path": f"assets/{asset.name}",
                "rights": asset.rights,
                "media_type": asset.media_type,
            }
            for asset in document.assets
        ],
        "handoff": "本地内容包；目标编辑器需要用户人工导入及预览，不证明平台兼容或发布",
        "resource_status": "complete",
        "platform_handoff_verified": False,
    }
    files["manifest.json"] = json.dumps(manifest, ensure_ascii=False, indent=2).encode(
        "utf-8"
    )
    stream = BytesIO()
    with ZipFile(stream, "w", compression=ZIP_DEFLATED) as archive:
        for name, data in files.items():
            archive.writestr(name, data)
    return _artifact(document, "zip", "application/zip", stream.getvalue())


def export_document(document: ExportDocument, format: ExportFormat) -> ExportArtifact:
    functions = {
        "md": export_markdown,
        "docx": export_docx,
        "html": export_html,
        "zip": export_zip,
    }
    if format not in functions:
        raise ValueError("不支持的文档导出格式")
    return functions[format](document)


def export_csv(
    rows: Sequence[Mapping[str, object]], fieldnames: Sequence[str]
) -> ExportArtifact:
    if not fieldnames or len(rows) > 20_000:
        raise ValueError("CSV 缺少字段或行数超过范围")
    stream = StringIO(newline="")
    writer = csv.DictWriter(stream, fieldnames=fieldnames, extrasaction="raise")
    writer.writeheader()
    for row in rows:
        values = {}
        for key, value in row.items():
            text = "" if value is None else str(value)
            # Keep visible content while preventing spreadsheet formula execution.
            values[key] = "'" + text if re.match(r"^[\s]*[=+@-]", text) else text
        writer.writerow(values)
    return ExportArtifact(
        "shots.csv", "text/csv; charset=utf-8", stream.getvalue().encode("utf-8-sig")
    )
