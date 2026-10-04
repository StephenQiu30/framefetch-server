"""Formal document tasks use isolated PostgreSQL, Temporal, and report publication."""

from __future__ import annotations

import asyncio
import hashlib
import json
from datetime import UTC, datetime, timedelta
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock
from uuid import uuid4

import pytest
from app.core.config import Settings
from app.integrations.analysis_report_docx import PythonDocxAnalysisReportRenderer
from app.integrations.analysis_skill_catalog import BuiltinAnalysisSkillCatalog
from app.integrations.temporal_client import CommandPublisher
from app.models import (
    AnalysisDocumentLockRow,
    AnalysisRunRow,
    AnalysisStepResultRow,
    DocumentArtifactRow,
    DocumentRow,
)
from app.repositories.analysis.execution import AnalysisExecutionPersistence
from app.repositories.analysis.report_repository import (
    SqlAlchemyAnalysisReportRepository,
)
from app.repositories.analysis.repository import SqlAlchemyAnalysisRepository
from app.repositories.documents.delete_repository import (
    SqlAlchemyDocumentDeleteRepository,
)
from app.repositories.downloads.repository import SqlAlchemyDownloadRepository
from app.repositories.outbox_repository import SqlAlchemyOutboxRepository
from app.services.analysis.create_document_analysis import CreateDocumentAnalysis
from app.services.analysis.errors import AnalysisApplicationError, PersistenceConflict
from app.services.analysis.models import AnalysisRetry
from app.services.analysis_execution.document_organization import (
    DocumentOrganizationExecutor,
    source_blocks,
)
from app.services.analysis_execution.errors import AnalysisOwnershipLost
from app.services.analysis_execution.ports import AnalyzerSelection
from app.services.analysis_execution.screenplay_executor import (
    ScreenplayAnalysisExecutor,
)
from app.services.documents.models import DocumentSnapshot, DocumentTextArtifactSnapshot
from app.services.documents.preview import read_document_preview
from app.services.downloads.fingerprints import HmacRequestFingerprinter
from app.services.imports.errors import ImportPersistenceConflict
from app.services.quotas import UserQuota
from app.workers.analysis.screenplay_artifacts import LocalScreenplayArtifactLoader
from app.workers.analysis.screenplay_providers import (
    ConfiguredScreenplayAnalyzerResolver,
)
from app.workers.analysis.skill_activities import SkillActivities
from app.workers.analysis.skill_workflow import SkillCommand, SkillWorkflow
from app.workers.outbox.loop import OutboxPublisherLoop
from app.workers.report.message import parse_report_requested
from app.workers.report.publisher import ReportPublisher
from docx import Document
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import async_sessionmaker
from temporalio.worker import Worker
from tests.unit.workers.analysis.fixtures import valid_screenplay_mapping

OWNER = "a" * 64
SOURCE = """# 已有观点

保留作者观点与限定。

```python
print("完整代码")
```

| 项 | 限定 |
| --- | --- |
| 一 | 仅本例 |

超过预览范围后仍应保留的末尾🙂。"""


class Storage:
    bucket = "video-artifacts"

    def __init__(self):
        self.objects = {}

    async def read_bounded(self, key, *, maximum):
        data = self.objects[key]
        if len(data) > maximum:
            raise ValueError("source limit")
        return data

    async def read_range(self, key, *, length):
        return self.objects[key][:length]

    async def download(self, key, target):
        target.write_bytes(self.objects[key])

    async def stat(self, key):
        raw = self.objects.get(key)
        return (
            None
            if raw is None
            else SimpleNamespace(
                size_bytes=len(raw), sha256=hashlib.sha256(raw).hexdigest()
            )
        )

    async def upload_bytes(self, key, content, media_type, sha):
        assert hashlib.sha256(content).hexdigest() == sha
        self.objects[key] = content


class Generator:
    def __init__(self, text):
        self.text = text
        self.requests = []

    async def generate_content(self, request):
        self.requests.append(request)
        bound = json.loads(request.prompt.split("\n\n")[-1])
        assert "".join(block["text"] for block in bound["blocks"]) == self.text
        blocks = source_blocks(self.text)
        return {
            "groups": [
                {
                    "block_ids": [block.id for block in blocks],
                    "heading_block_id": blocks[0].id if blocks[0].heading else None,
                    "reason": (
                        "原文先提出观点，再用代码和表格说明具体边界；保留结"
                        "尾限定，当前未提供另一个事实来源。"
                    ),
                }
            ]
        }

    async def analyze_screenplay(self, request):
        self.requests.append(request)
        assert request.screenplay_text == self.text
        assert request.source_index_json
        value = valid_screenplay_mapping()
        value["scenes"] = [
            dict(value["scenes"][0], id=f"note-{index}", source_scene_id=identifier)
            for index, identifier in enumerate(request.source_scene_ids)
        ]
        return value

    async def synthesize_screenplay_analysis(self, request):
        raise AssertionError("short source is one call")


class Resolver:
    def __init__(self, generator):
        self.generator = generator

    async def resolve(self):
        return AnalyzerSelection(
            self.generator, "controlled", "controlled", "controlled", "b" * 64
        )


async def seed_document(sessions, storage, text=SOURCE):
    now = datetime.now(UTC)
    identifier = uuid4()
    raw = text.encode()
    digest = hashlib.sha256(raw).hexdigest()
    key = f"documents/{identifier}/1/document.md"
    storage.objects[key] = raw
    async with sessions() as session, session.begin():
        session.add(
            DocumentRow(
                id=identifier,
                owner_hash=OWNER,
                idempotency_key=str(uuid4()),
                request_fingerprint="f" * 64,
                title="源文件名不是作者标题",
                original_filename="ordinary.md",
                source_format="markdown",
                content_type="text/markdown",
                declared_size_bytes=len(raw),
                declared_sha256=digest,
                rights_statement_version="v1",
                status="ready",
                attempt=1,
                detected_language="zh-CN",
                scene_count=0,
                character_count=len(text),
                text_sha256=digest,
                finished_at=now,
                created_at=now,
                updated_at=now,
            )
        )
        await session.flush()
        session.add(
            DocumentArtifactRow(
                id=uuid4(),
                document_id=identifier,
                kind="original",
                bucket="video-artifacts",
                object_key=f"documents/{identifier}/1/original.md",
                content_type="text/markdown",
                size_bytes=len(raw),
                sha256=digest,
                status="ready",
                created_at=now,
                updated_at=now,
            )
        )
        session.add(
            DocumentArtifactRow(
                id=uuid4(),
                document_id=identifier,
                kind="normalized",
                bucket="video-artifacts",
                object_key=key,
                content_type="text/markdown; charset=utf-8",
                size_bytes=len(raw),
                sha256=digest,
                status="ready",
                artifact_metadata={"scenes": [], "parse_summary": {}},
                created_at=now,
                updated_at=now,
            )
        )
    return identifier, digest


class Fallback:
    def __init__(self):
        self.events = []

    async def publish(self, envelope):
        self.events.append(envelope)


def creator(repository):
    return CreateDocumentAnalysis(
        repository=repository,
        fingerprinter=HmacRequestFingerprinter(b"controlled-key-16"),
        now=lambda: datetime.now(UTC),
        new_id=uuid4,
        max_attempts=1,
        skill_catalog=BuiltinAnalysisSkillCatalog(),
        enabled=True,
    )


@pytest.mark.parametrize("skill_id", ["article-format", "wechat-format", "xhs-format"])
async def test_formal_document_outbox_temporal_full_source_reports_and_owner_guards(
    postgres_engine, temporal_client, tmp_path, skill_id
):
    sessions = async_sessionmaker(postgres_engine, expire_on_commit=False)
    repository = SqlAlchemyAnalysisRepository(sessions)
    storage = Storage()
    identifier, digest = await seed_document(sessions, storage)
    # Controlled 24-character preview proves execution reads the owned full object.
    # It does not claim this small fixture exceeds the production preview limit.
    now = datetime.now(UTC)
    preview, truncated = await read_document_preview(
        DocumentSnapshot(
            id=identifier,
            owner_hash=OWNER,
            title="ordinary.md",
            original_filename="ordinary.md",
            source_format="markdown",
            declared_size_bytes=len(SOURCE.encode()),
            status="ready",
            attempt=1,
            error_code=None,
            version=1,
            detected_language="zh-CN",
            scene_count=0,
            character_count=len(SOURCE),
            text_sha256=digest,
            quality_warnings=(),
            parse_summary=None,
            created_at=now,
            updated_at=now,
            finished_at=now,
            normalized_artifact=DocumentTextArtifactSnapshot(
                storage.bucket,
                f"documents/{identifier}/1/document.md",
                len(SOURCE.encode()),
                digest,
            ),
        ),
        storage,
        max_bytes=128,
        max_characters=24,
    )
    assert truncated and preview == SOURCE[:24]
    assert "末尾🙂" not in preview
    create = creator(repository)
    with pytest.raises(AnalysisApplicationError):
        await create(identifier, "c" * 64, "foreign", "article-format", "zh-CN")
    with pytest.raises(AnalysisApplicationError):
        await create(identifier, OWNER, "inactive", "screenplay-rewrite", "zh-CN")
    with pytest.raises(AnalysisApplicationError):
        await create(
            identifier, OWNER, "too-long", "article-format", "zh-CN", "x" * 4001
        )
    view = await create(
        identifier,
        OWNER,
        "fixed-request",
        skill_id,
        "en-US",
        quota=UserQuota(exempt=True),
    )
    replay = await create(
        identifier,
        OWNER,
        "fixed-request",
        skill_id,
        "en-US",
        quota=UserQuota(exempt=True),
    )
    assert replay.id == view.id
    with pytest.raises(AnalysisApplicationError):
        await create(
            identifier,
            OWNER,
            "fixed-request",
            skill_id,
            "zh-CN",
            quota=UserQuota(exempt=True),
        )
    with pytest.raises(ImportPersistenceConflict):
        await SqlAlchemyDocumentDeleteRepository(sessions).prepare_document_deletion(
            identifier, OWNER, now=datetime.now(UTC)
        )
    generator = Generator(SOURCE)
    resolver = Resolver(generator)
    persistence = AnalysisExecutionPersistence(
        repository, SqlAlchemyDownloadRepository(sessions)
    )
    executor = DocumentOrganizationExecutor(
        repository,
        storage,
        resolver,
        tmp_path,
        bucket="video-artifacts",
        maximum_bytes=2 * 1024**2,
        maximum_characters=120000,
        timeout_seconds=900,
    )
    activities = SkillActivities(
        repository,
        persistence,
        {("screenplay", "structured-report"): executor},
        Settings(app_env="test", _env_file=None),
    )
    fallback = Fallback()
    publisher = CommandPublisher(
        fallback,
        object(),
        address="unused",
        namespace="framefetch-test",
        cancel_inspection=AsyncMock(),
        analyses=repository,
    )
    publisher._client = temporal_client
    loop = OutboxPublisherLoop(
        repository=SqlAlchemyOutboxRepository(sessions),
        publisher=publisher,
        publisher_id="test-formal",
        clock=lambda: datetime.now(UTC),
    )
    command = SkillCommand(str(view.id), str(view.run_id), view.run_no)
    async with Worker(
        temporal_client,
        task_queue="ff-skill",
        workflows=[SkillWorkflow],
        activities=[activities.run, activities.reconcile],
    ):
        assert await loop.run_once() >= 1
        status = await asyncio.wait_for(
            temporal_client.get_workflow_handle(command.workflow_id).result(), 30
        )
        snapshot = await repository.get_job(view.id)
        assert status == "running", (
            status,
            snapshot.error_code,
            len(generator.requests),
        )
    assert len(generator.requests) == 1
    await loop.run_once()
    report = [
        item
        for item in fallback.events
        if item.event_type == "analysis.report.publish.requested"
    ]
    assert len(report) == 1
    report_publisher = ReportPublisher(
        SqlAlchemyAnalysisReportRepository(sessions),
        storage,
        PythonDocxAnalysisReportRenderer(),
        bucket="video-artifacts",
        worker_id="test-report",
        clock=lambda: datetime.now(UTC),
    )
    assert await report_publisher.execute(parse_report_requested(report[0].to_bytes()))
    current = await repository.get_job(view.id)
    assert current.status == "succeeded"
    result = await repository.get_result(view.id)
    assert result.media is None
    result.validate_document_source(digest)
    markdown = await repository.get_current_report_file(view.id, "markdown")
    assert SOURCE.encode() in storage.objects[markdown.object_key]
    docx = await repository.get_current_report_file(view.id, "docx")
    document = Document(BytesIO(storage.objects[docx.object_key]))
    assert "完整代码" in "\n".join(p.text for p in document.paragraphs)
    assert len(document.tables) == 1
    async with sessions() as session:
        run = await session.get(AnalysisRunRow, view.run_id)
        assert (
            run.model_calls_used == 1 and run.execution_binding["max_model_calls"] == 1
        )
        assert run.execution_binding["source_sha256"] == digest
        assert run.execution_deadline <= run.started_at + timedelta(seconds=901)
        assert (
            await session.scalar(
                select(func.count()).select_from(AnalysisStepResultRow)
            )
            == 1
        )
    # Deleting the actual analysis releases its ordinary source lock.
    await repository.delete_job(view.id, OWNER, datetime.now(UTC))
    async with sessions() as session:
        assert (
            await session.scalar(
                select(func.count()).select_from(AnalysisDocumentLockRow)
            )
            == 0
        )


async def test_zero_scene_novel_executes_one_call_from_exact_persisted_text(
    postgres_engine, tmp_path, monkeypatch
):
    from app.workers.analysis import skill_activities

    monkeypatch.setattr(skill_activities.activity, "heartbeat", lambda: None)
    sessions = async_sessionmaker(postgres_engine, expire_on_commit=False)
    repository = SqlAlchemyAnalysisRepository(sessions)
    storage = Storage()
    text = (
        "第一章 取件\n林澈想取父亲的信。\n\n第二章 共同拆"
        "信\n林遥说兄妹需共同写地址。é🙂"
    )
    identifier, digest = await seed_document(sessions, storage, text)
    view = await creator(repository)(
        identifier,
        OWNER,
        "novel",
        "screenplay-analysis",
        "zh-CN",
        quota=UserQuota(exempt=True),
    )
    generator = Generator(text)
    resolver = Resolver(generator)
    loader = LocalScreenplayArtifactLoader(
        storage,
        workspace_root=tmp_path,
        bucket="video-artifacts",
        max_source_bytes=2 * 1024**2,
    )
    persistence = AnalysisExecutionPersistence(
        repository, SqlAlchemyDownloadRepository(sessions)
    )
    executor = ScreenplayAnalysisExecutor(
        repository=persistence,
        loader=loader,
        resolver=ConfiguredScreenplayAnalyzerResolver(resolver),
        clock=lambda: datetime.now(UTC),
        max_single_call_characters=120000,
    )
    activities = SkillActivities(
        repository,
        persistence,
        {("screenplay", "screenplay-analysis"): executor},
        Settings(app_env="test", _env_file=None),
    )
    assert (
        await activities.run(SkillCommand(str(view.id), str(view.run_id), view.run_no))
        == "running"
    )
    assert len(generator.requests) == 1
    result = await repository.get_result(view.id)
    assert [item.source_scene_id for item in result.scenes] == ["unit-1", "unit-2"]
    async with sessions() as session:
        run = await session.get(AnalysisRunRow, view.run_id)
        assert run.execution_binding["source_sha256"] == digest
        assert run.execution_binding["source_units"][-1]["end"] == len(text)
        assert run.model_calls_used == 1
    assert not list(tmp_path.glob("analysis-*"))


async def test_unknown_step_blocks_atomic_retry_and_late_receipt(postgres_engine):
    sessions = async_sessionmaker(postgres_engine, expire_on_commit=False)
    repository = SqlAlchemyAnalysisRepository(sessions)
    storage = Storage()
    identifier, _ = await seed_document(sessions, storage)
    view = await creator(repository)(
        identifier,
        OWNER,
        "unknown",
        "article-format",
        "zh-CN",
        quota=UserQuota(exempt=True),
    )
    now = datetime.now(UTC)
    worker = "test-owner"
    job = await repository.claim_run(
        view.id, view.run_id, view.run_no, worker, now, timedelta(seconds=60)
    )
    await repository.bind_execution(
        view.run_id,
        {"max_model_calls": 1},
        owner=worker,
        attempt=job.attempt,
        now=now,
        deadline=now + timedelta(seconds=900),
    )
    await repository.begin_step(
        view.run_id, "plan", "d" * 64, owner=worker, attempt=job.attempt, now=now
    )
    terminal = await repository.fail_run(
        view.id, view.run_id, error_code="worker_lost", now=datetime.now(UTC)
    )
    assert terminal.error_code == "analysis_outcome_unknown"
    with pytest.raises(AnalysisOwnershipLost):
        await repository.complete_step(
            view.run_id,
            "plan",
            {"late": True},
            owner=worker,
            attempt=job.attempt,
            now=datetime.now(UTC),
        )
    retry = AnalysisRetry(
        job_id=view.id,
        run_id=uuid4(),
        owner_hash=OWNER,
        idempotency_key="retry",
        trigger="manual_retry",
        max_attempts=1,
        outbox_event_id=uuid4(),
        max_runs_per_job=3,
        min_interval_seconds=0,
        retries_per_day=10,
        quota=UserQuota(exempt=True),
    )
    with pytest.raises(PersistenceConflict):
        await repository.retry_job_and_enqueue(retry, now=datetime.now(UTC))
    async with sessions() as session:
        assert (
            await session.scalar(select(func.count()).select_from(AnalysisRunRow)) == 1
        )
        assert (await session.get(AnalysisRunRow, view.run_id)).model_calls_used == 1
