"""Real API, isolated PostgreSQL and Temporal; deterministic work calls no model."""

import asyncio
import base64
from datetime import timedelta
from io import BytesIO
from pathlib import Path
from random import Random
from types import SimpleNamespace
from unittest.mock import AsyncMock
from uuid import UUID, uuid4

import pytest
from app.api.deps import (
    get_creation_export_service,
    get_creation_service,
    get_current_user,
    get_runtime_settings,
)
from app.api.errors import register_exception_handlers
from app.api.routes.creation import router
from app.api.routes.creation_exports import router as export_router
from app.core.config import Settings
from app.core.db import utc_now
from app.core.errors import AppError
from app.integrations.temporal_client import CommandPublisher
from app.main import create_app
from app.models.analysis_worker import AnalysisWorkerHeartbeatRow
from app.models.creation import CreationExportRow, CreationRevisionRow
from app.models.outbox import OutboxEventRow
from app.repositories.analysis.worker_registry import ANALYSIS_MESSAGE_SCHEMA_VERSION
from app.repositories.creation import CreationRepository
from app.repositories.outbox_repository import SqlAlchemyOutboxRepository
from app.services.analysis_execution.ports import AnalyzerSelection
from app.services.auth.models import CurrentUser, UserRole
from app.services.creation.catalog import get_capability
from app.services.creation.export_service import CreationExportService
from app.services.creation.models import (
    CreationConfirmRequest,
    CreationTaskCreateRequest,
)
from app.services.creation.service import CreationService
from app.workers.analysis.creation_activities import CreationActivities
from app.workers.analysis.creation_workflow import CreationCommand, CreationWorkflow
from app.workers.outbox.loop import OutboxPublisherLoop
from docx import Document
from fastapi import FastAPI
from httpx import ASGITransport, AsyncClient
from PIL import Image
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import async_sessionmaker
from temporalio.worker import Replayer, Worker


class NeverResolve:
    def __init__(self):
        self.calls = 0

    async def resolve(self):
        self.calls += 1
        raise AssertionError("deterministic preparation must not resolve a model")


def api(repository):
    service = CreationService(repository)
    exports = CreationExportService(repository, font_path=None)
    now = utc_now()
    user = CurrentUser(uuid4(), "test", "test@example.test", UserRole.ADMIN, now, now)
    application = FastAPI()
    application.state.services = SimpleNamespace()
    application.include_router(router, prefix="/api")
    application.include_router(export_router, prefix="/api")
    application.dependency_overrides[get_current_user] = lambda: user
    application.dependency_overrides[get_creation_service] = lambda: service
    application.dependency_overrides[get_creation_export_service] = lambda: exports
    application.dependency_overrides[get_runtime_settings] = lambda: Settings(
        _env_file=None
    )
    register_exception_handlers(application)
    return application, user


async def ready(sessions):
    async with sessions() as session, session.begin():
        session.add(
            AnalysisWorkerHeartbeatRow(
                worker_id=str(uuid4()),
                app_version="test",
                message_schema_version=ANALYSIS_MESSAGE_SCHEMA_VERSION,
                last_seen_at=utc_now(),
            )
        )


async def test_api_outbox_worker_revision_and_file_export_are_one_immutable_flow(
    postgres_engine, temporal_client, tmp_path: Path
):
    sessions = async_sessionmaker(postgres_engine, expire_on_commit=False)
    repository = CreationRepository(sessions)
    await ready(sessions)
    application, user = api(repository)
    resolver = NeverResolve()
    activities = CreationActivities(
        repository,
        resolver,
        AsyncMock(),
        Settings(_env_file=None, analysis_workspace_root=tmp_path),
    )
    queue = "ff-skill"
    publisher = CommandPublisher(
        AsyncMock(),
        AsyncMock(),
        address=temporal_client.service_client.config.target_host,
        namespace="framefetch-test",
        cancel_inspection=AsyncMock(),
        creation=repository,
    )
    loop = OutboxPublisherLoop(
        repository=SqlAlchemyOutboxRepository(sessions),
        publisher=publisher,
        publisher_id=str(uuid4()),
        clock=utc_now,
    )
    async with AsyncClient(
        transport=ASGITransport(app=application), base_url="http://test"
    ) as client:
        created = await client.post(
            "/api/creation/materials",
            headers={"Idempotency-Key": "material"},
            json={
                "kind": "text",
                "title": "原稿",
                "text": "# 验证原稿\n\n这是一段保留的正文。",
                "rights_statement": "本人原创",
            },
        )
        assert created.status_code == 201
        original = created.json()["data"]
        material_id = original["id"]
        revision_id = original["current_revision"]["id"]
        confirmed = await client.post(
            f"/api/creation/materials/{material_id}/confirm",
            json={"expected_revision_id": revision_id},
        )
        assert confirmed.status_code == 200
        body = {
            "skill_id": "article-edit",
            "material_revision_ids": [revision_id],
            "options": {"mode": "format"},
        }
        first = await client.post(
            "/api/creation/tasks", headers={"Idempotency-Key": "create-task"}, json=body
        )
        assert first.status_code == 201
        task_id = first.json()["data"]["id"]
        assert first.json()["data"]["status"] == "queued"
        replay = await client.post(
            "/api/creation/tasks", headers={"Idempotency-Key": "create-task"}, json=body
        )
        assert replay.json()["data"]["id"] == task_id
        command = CreationCommand(task_id, 1)
        assert await loop.run_once() == 1
        handle = temporal_client.get_workflow_handle(command.workflow_id)
        run_id = (await handle.describe()).run_id
        async with sessions() as session, session.begin():
            event = await session.scalar(
                select(OutboxEventRow).where(
                    OutboxEventRow.aggregate_id == UUID(task_id)
                )
            )
            event.published_at = None
        assert await loop.run_once() == 1
        assert (await handle.describe()).run_id == run_id
        async with Worker(
            temporal_client,
            task_queue=queue,
            workflows=[CreationWorkflow],
            activities=[activities.run, activities.reconcile],
            graceful_shutdown_timeout=timedelta(seconds=1),
        ):
            assert (
                await asyncio.wait_for(handle.result(), 30) == "awaiting_confirmation"
            )
        assert resolver.calls == 0
        result = (await client.get(f"/api/creation/tasks/{task_id}")).json()["data"]
        candidate_id = result["revision"]["id"]
        assert result["usage"]["calls_used"] == result["usage"]["calls_reserved"] == 0
        premature = await client.get(
            f"/api/creation/tasks/{task_id}/revisions/{candidate_id}/export/md"
        )
        assert premature.status_code == 409
        assert (
            await client.post(
                f"/api/creation/tasks/{task_id}/confirm",
                json={"expected_revision_id": candidate_id},
            )
        ).status_code == 200
        path = f"/api/creation/tasks/{task_id}/revisions/{candidate_id}/export/md"
        exported = await client.get(path)
        assert exported.status_code == 200
        assert "这是一段保留的正文。".encode() in exported.content
        assert exported.headers["cache-control"] == "private, no-store"
        assert (await client.get(path)).content == exported.content
        edited = await client.post(
            f"/api/creation/tasks/{task_id}/revisions",
            headers={"Idempotency-Key": "manual-edit"},
            json={
                "expected_revision_id": candidate_id,
                "text": "人工第二版",
                "data": {},
            },
        )
        assert edited.status_code == 200
        assert (await client.get(path)).content == exported.content
        conflict = await client.post(
            f"/api/creation/tasks/{task_id}/revisions",
            headers={"Idempotency-Key": "stale-edit"},
            json={"expected_revision_id": candidate_id, "text": "旧页覆盖", "data": {}},
        )
        assert conflict.status_code == 409
        source = await client.get(f"/api/creation/materials/{material_id}/revisions")
        assert source.json()["data"][0]["id"] == revision_id
        another_user = CurrentUser(
            uuid4(),
            "foreign",
            "foreign@example.test",
            UserRole.ADMIN,
            utc_now(),
            utc_now(),
        )
        application.dependency_overrides[get_current_user] = lambda: another_user
        assert (await client.get(f"/api/creation/tasks/{task_id}")).status_code == 404
        assert (await client.get(path)).status_code == 404
        assert (await client.get("/api/creation/materials")).json()["data"] == []
    async with sessions() as session:
        assert (
            await session.scalar(select(func.count()).select_from(CreationExportRow))
            == 1
        )
        assert (
            await session.scalar(select(func.count()).select_from(CreationRevisionRow))
            == 3
        )
    history = await handle.fetch_history()
    await Replayer(workflows=[CreationWorkflow]).replay_workflow(history)
    assert "这是一段保留的正文。".encode() not in b"".join(
        event.SerializeToString() for event in history.events
    )
    assert (
        await repository.get_task(UUID(task_id), user.owner_hash)
    ).usage.calls_used == 0


async def test_api_rejects_unknown_cost_retry_and_foreign_materials(postgres_engine):
    sessions = async_sessionmaker(postgres_engine, expire_on_commit=False)
    repository = CreationRepository(sessions)
    await ready(sessions)
    application, user = api(repository)
    async with AsyncClient(
        transport=ASGITransport(app=application), base_url="http://test"
    ) as client:
        material = (
            await client.post(
                "/api/creation/materials",
                headers={"Idempotency-Key": "material"},
                json={
                    "kind": "text",
                    "title": "原稿",
                    "text": "原稿正文",
                    "rights_statement": "原创",
                },
            )
        ).json()["data"]
        await repository.confirm_material(
            UUID(material["id"]),
            user.owner_hash,
            CreationConfirmRequest(
                expected_revision_id=UUID(material["current_revision"]["id"])
            ),
        )
        value = await repository.create_task(
            user.owner_hash,
            "direct-fixture",
            CreationTaskCreateRequest(
                skill_id="article-edit",
                material_revision_ids=(UUID(material["current_revision"]["id"]),),
            ),
            "c" * 64,
        )
        await repository.claim_task(value.id, "fixture", expected_attempt=1)
        await repository.mark_call_sent(value.id, "fixture")
        await repository.fail_task(
            value.id, "fixture", "worker_lost", outcome_unknown=True
        )
        retry = await client.post(
            f"/api/creation/tasks/{value.id}/attempts",
            headers={"Idempotency-Key": "unsafe-retry"},
            json={"acknowledge_unknown_cost": True},
        )
        assert retry.status_code == 409
        state = (await client.get(f"/api/creation/tasks/{value.id}")).json()["data"]
        assert (
            state["usage"]["calls_reserved"]
            == state["usage"]["unknown_operations"]
            == 1
        )
        assert state["attempt"] == 1
        malformed = await client.post(
            "/api/creation/materials",
            headers={"Idempotency-Key": "bad-image"},
            json={
                "kind": "image",
                "title": "假图片",
                "image_data_base64": "dGV4dA==",
                "rights_statement": "原创",
            },
        )
        assert malformed.status_code == 422
        duplicate_body = {
            "skill_id": "article-edit",
            "material_revision_ids": [material["current_revision"]["id"]],
            "unexpected": True,
        }
        assert (
            await client.post(
                "/api/creation/tasks",
                headers={"Idempotency-Key": "bad-options"},
                json=duplicate_body,
            )
        ).status_code == 422


async def test_uploaded_document_keeps_original_and_exact_history(postgres_engine):
    sessions = async_sessionmaker(postgres_engine, expire_on_commit=False)
    repository = CreationRepository(sessions)
    application, _ = api(repository)
    original = "# 真实上传\n\n作者提供的正文。".encode()
    body = {
        "kind": "text",
        "title": "上传原稿",
        "rights_statement": "本人原创",
        "document_filename": "original.md",
        "document_data_base64": base64.b64encode(original).decode(),
    }
    async with AsyncClient(
        transport=ASGITransport(app=application), base_url="http://test"
    ) as client:
        first = await client.post(
            "/api/creation/materials", headers={"Idempotency-Key": "doc"}, json=body
        )
        assert first.status_code == 201
        material = first.json()["data"]
        assert material["current_revision"]["text"] == original.decode()
        path = f"/api/creation/materials/{material['id']}/source"
        binary = await client.get(path)
        assert binary.status_code == 200 and binary.content == original
        assert binary.headers["cache-control"] == "private, no-store"
        assert "original.md" in binary.headers["content-disposition"]
        assert (
            await client.get(f"/api/creation/materials/{material['id']}/image")
        ).status_code == 404
        saved = await client.post(
            f"/api/creation/materials/{material['id']}/revisions",
            headers={"Idempotency-Key": "manual"},
            json={
                "expected_revision_id": material["current_revision"]["id"],
                "text": "人工整理正文",
                "data": {"file_sha256": "false"},
            },
        )
        assert saved.status_code == 200
        assert (
            saved.json()["data"]["current_revision"]["data"]["file_sha256"]
            == material["current_revision"]["data"]["file_sha256"]
        )
        assert (await client.get(path)).content == original
        mismatch = await client.post(
            "/api/creation/materials",
            headers={"Idempotency-Key": "false"},
            json={**body, "text": "虚挂来源正文"},
        )
        assert mismatch.status_code == 422
        foreign = CurrentUser(
            uuid4(),
            "foreign",
            "foreign@example.test",
            UserRole.ADMIN,
            utc_now(),
            utc_now(),
        )
        application.dependency_overrides[get_current_user] = lambda: foreign
        assert (await client.get(path)).status_code == 404


async def test_real_request_guard_accepts_bounded_original_and_revision_preview(
    postgres_engine,
):
    sessions = async_sessionmaker(postgres_engine, expire_on_commit=False)
    repository = CreationRepository(sessions)
    prepared, _ = api(repository)
    application = create_app(Settings(app_env="test", _env_file=None))
    application.dependency_overrides.update(prepared.dependency_overrides)
    image_buffer = BytesIO()
    Image.frombytes("RGB", (512, 512), Random(3).randbytes(512 * 512 * 3)).save(
        image_buffer, format="PNG"
    )
    document = Document()
    document.add_paragraph("实际DOCX正文")
    image_buffer.seek(0)
    document.add_picture(image_buffer)
    source = BytesIO()
    document.save(source)
    original = source.getvalue()
    assert 256 * 1024 < len(original) < 10 * 1024 * 1024
    async with AsyncClient(
        transport=ASGITransport(app=application),
        base_url="http://test",
        headers={"Origin": "http://test"},
    ) as client:
        response = await client.post(
            "/api/creation/materials",
            headers={"Idempotency-Key": "guard-doc"},
            json={
                "kind": "text",
                "title": "真实文档",
                "rights_statement": "本人原创",
                "document_filename": "original.docx",
                "document_data_base64": base64.b64encode(original).decode(),
            },
        )
        assert response.status_code == 201
        material = response.json()["data"]
        assert material["current_revision"]["text"] == "实际DOCX正文\n\n"
        assert (
            await client.get(f"/api/creation/materials/{material['id']}/source")
        ).content == original
        saved = await client.post(
            f"/api/creation/materials/{material['id']}/revisions",
            headers={"Idempotency-Key": "guard-save"},
            json={
                "expected_revision_id": material["current_revision"]["id"],
                "text": "人工整理",
                "data": {"preview_fixture": "x" * 400_000},
            },
        )
        assert saved.status_code == 200
        unrelated = await client.post(
            "/api/auth/login", content=b"x" * (256 * 1024 + 1)
        )
        assert unrelated.status_code == 413
        malformed_path = await client.post(
            "/api/creation/materials/not-a-uuid/revisions",
            content=b"x" * (256 * 1024 + 1),
        )
        assert malformed_path.status_code == 413
        declared = await client.post(
            "/api/creation/materials",
            headers={"Content-Length": str(16 * 1024 * 1024 + 1)},
            content=b"{}",
        )
        assert declared.status_code == 413
        streamed = await client.post(
            f"/api/creation/materials/{material['id']}/revisions",
            content=b"x" * (2 * 1024 * 1024 + 1),
        )
        assert streamed.status_code == 413


class LostReceiptGenerator:
    def __init__(self):
        self.calls = 0

    async def generate_content(self, request):
        self.calls += 1
        raise RuntimeError("transport disconnected after the provider accepted work")


class ControlledResolver:
    def __init__(self, analyzer):
        self.analyzer = analyzer

    async def resolve(self):
        return AnalyzerSelection(
            self.analyzer, "controlled", "controlled", "controlled", "a" * 64
        )


async def test_creation_workflow_retains_lost_call_and_never_auto_repeats(
    postgres_engine, temporal_client, tmp_path: Path
):
    sessions = async_sessionmaker(postgres_engine, expire_on_commit=False)
    repository = CreationRepository(sessions)
    application, user = api(repository)
    await ready(sessions)
    async with AsyncClient(
        transport=ASGITransport(app=application), base_url="http://test"
    ) as client:
        response = await client.post(
            "/api/creation/materials",
            headers={"Idempotency-Key": "source"},
            json={
                "kind": "text",
                "title": "授权材料",
                "text": "确实提供的写作事实",
                "rights_statement": "本人原创",
            },
        )
        material = response.json()["data"]
    await repository.confirm_material(
        UUID(material["id"]),
        user.owner_hash,
        CreationConfirmRequest(
            expected_revision_id=UUID(material["current_revision"]["id"])
        ),
    )
    value = await repository.create_task(
        user.owner_hash,
        "model-fixture",
        CreationTaskCreateRequest(
            skill_id="article-write",
            material_revision_ids=(UUID(material["current_revision"]["id"]),),
        ),
        get_capability("article-write").method_sha256,
    )
    analyzer = LostReceiptGenerator()
    activities = CreationActivities(
        repository,
        ControlledResolver(analyzer),
        AsyncMock(),
        Settings(_env_file=None, analysis_workspace_root=tmp_path),
    )
    command = CreationCommand(str(value.id), 1)
    async with Worker(
        temporal_client,
        task_queue="ff-skill",
        workflows=[CreationWorkflow],
        activities=[activities.run, activities.reconcile],
    ):
        result = await asyncio.wait_for(
            temporal_client.execute_workflow(
                CreationWorkflow.run,
                command,
                id=command.workflow_id,
                task_queue="ff-skill",
            ),
            30,
        )
    assert result == "outcome_unknown"
    assert analyzer.calls == 1
    state = await repository.get_task(value.id, user.owner_hash)
    assert state.usage.calls_used == 0
    assert state.usage.calls_reserved == state.usage.unknown_operations == 1
    assert state.usage.tokens_reserved == state.budget.max_tokens
    with pytest.raises(AppError) as unsafe:
        await repository.retry_task(value.id, user.owner_hash, "retry-without-receipt")
    assert unsafe.value.status == 409
    assert (
        await repository.claim_task(value.id, command.workflow_id, expected_attempt=1)
        is None
    )
    assert analyzer.calls == 1
