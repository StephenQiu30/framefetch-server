from io import BytesIO
from zipfile import ZipFile, is_zipfile

from app.integrations.analysis_report_docx import PythonDocxAnalysisReportRenderer
from app.services.analysis.report import render_analysis_report_markdown
from app.services.analysis.rules.enums import AnalysisResultKind
from docx import Document
from tests.unit.services.analysis.test_editorial_examples import _parse, _read
from tests.unit.services.analysis.test_report import report_result


def test_docx_report_is_valid_and_uses_business_brief_geometry() -> None:
    markdown = render_analysis_report_markdown(report_result())
    content = PythonDocxAnalysisReportRenderer().render(
        markdown, result_kind=AnalysisResultKind.VIDEO_VISUAL_ANALYSIS
    )

    assert content.startswith(b"PK")
    assert is_zipfile(BytesIO(content))
    document = Document(BytesIO(content))
    section = document.sections[0]
    assert round(section.page_width.inches, 2) == 8.5
    assert round(section.page_height.inches, 2) == 11
    assert round(section.left_margin.inches, 2) == 1
    assert document.core_properties.title == "产品 [演示](https://invalid.example)"
    assert document.core_properties.subject == "Editorial analysis report"
    assert len(document.tables) == 1
    paragraphs = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "核心判断" in paragraphs
    assert "修改建议" in paragraphs

    with ZipFile(BytesIO(content)) as package:
        xml = package.read("word/document.xml").decode("utf-8")
        header_xml = package.read("word/header1.xml").decode("utf-8")
    assert 'w:w="9360"' in xml
    assert 'w:w="2040"' in xml
    assert 'w:w="120"' in xml
    assert 'w:fill="F2F4F7"' in xml
    assert "ANALYSIS REPORT" in header_xml
    assert "EDITORIAL REVIEW" in header_xml
    assert "展示主要界面。" in xml


def test_article_docx_contains_only_article_copy_and_preserves_qualification() -> None:
    result = _parse("video-article", _read("video-article"))
    markdown = render_analysis_report_markdown(result)
    content = PythonDocxAnalysisReportRenderer().render(
        markdown, result_kind=result.kind
    )

    document = Document(BytesIO(content))
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    assert paragraphs[0] == result.title
    assert paragraphs[-1] == result.closing
    assert "还需要更完整的测试" in "\n".join(paragraphs)
    assert document.core_properties.subject == "Article"
    assert document.core_properties.author == ""
    with ZipFile(BytesIO(content)) as package:
        assert not any(
            name.startswith(("word/header", "word/footer"))
            for name in package.namelist()
        )
        xml = package.read("word/document.xml").decode("utf-8")
    for metadata in ("编辑摘要", "编辑附录", "00:07.000", "ANALYSIS REPORT"):
        assert metadata not in xml
