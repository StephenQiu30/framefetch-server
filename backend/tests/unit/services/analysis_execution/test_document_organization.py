from __future__ import annotations

import hashlib
from dataclasses import replace
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock

import pytest
from app.integrations.analysis_report_docx import PythonDocxAnalysisReportRenderer
from app.integrations.analysis_skill_catalog import BuiltinAnalysisSkillCatalog
from app.repositories.analysis.repository_serialization import (
    analysis_result_document,
    analysis_result_from_document,
)
from app.services.analysis.report import render_analysis_report_markdown
from app.services.analysis.rules.enums import AnalysisInputKind
from app.services.analysis_execution.document_organization import (
    document_from_plan,
    source_blocks,
)
from app.services.analysis_execution.errors import AnalysisArtifactError
from app.services.analysis_execution.models import ScreenplaySceneSource
from app.services.analysis_execution.screenplay_executor import (
    ScreenplayAnalysisExecutor,
)
from app.services.analysis_execution.screenplay_source_validation import (
    validate_screenplay_source,
)
from app.services.skills.drama import index_drama_source
from docx import Document
from tests.unit.workers.analysis.fakes import NOW, running_job

SOURCE = """# 原有标题

作者主张：保留全文🙂。

```python
print("不能拆开的代码")

print("尾行")
```

- 第一项
  - 嵌套项
- 第二项

| 项目 | 条件 |
| --- | --- |
| A | 仅限本例 |

[已有链接][资料]

[资料]: https://example.com/source

末尾不可省略的限定：未经验证，不保证效果。"""


def plan(blocks):
    return {
        "groups": [
            {
                "block_ids": [block.id for block in blocks],
                "heading_block_id": blocks[0].id,
                "reason": (
                    "先提出保留原文的主张，再用代码、列表和表格说明处理"
                    "边界；末尾限定需保留。"
                ),
            }
        ]
    }


@pytest.mark.parametrize("skill_id", ["article-format", "wechat-format", "xhs-format"])
def test_full_markdown_protected_blocks_survive_report_docx_and_storage(skill_id):
    blocks = source_blocks(SOURCE)
    assert "".join(block.text for block in blocks) == SOURCE
    for literal in ["```python\nprint(", "- 第一项\n  - 嵌套项", "| 项目 | 条件 |"]:
        assert sum(literal in block.text for block in blocks) == 1
    digest = hashlib.sha256(SOURCE.encode()).hexdigest()
    result = document_from_plan(
        plan(blocks), blocks, sha256=digest, language="zh-CN", skill_id=skill_id
    )
    assert result.media is None
    assert analysis_result_from_document(analysis_result_document(result)) == result
    markdown = render_analysis_report_markdown(result)
    assert SOURCE in markdown
    assert markdown.count("# 原有标题") == 1
    assert markdown.count(digest) == 1
    assert "原文分组" not in markdown
    assert f"原有标题 · Unicode \\[{0}, {len(SOURCE)}\\)" in markdown
    docx = PythonDocxAnalysisReportRenderer().render(
        markdown, result_kind="structured_report"
    )
    document = Document(BytesIO(docx))
    visible = "\n".join(p.text for p in document.paragraphs)
    assert 'print("不能拆开的代码")' in visible
    assert "末尾不可省略的限定：未经验证，不保证效果。" in visible
    assert len(document.tables) == 1
    assert document.tables[0].rows[1].cells[1].text == "仅限本例"


@pytest.mark.parametrize(
    "mutation",
    ["omit-tail", "duplicate", "reorder", "invent-heading", "foreign-heading"],
)
def test_organization_rejects_loss_order_changes_and_fabricated_headings(mutation):
    blocks = source_blocks(SOURCE)
    value = plan(blocks)
    group = value["groups"][0]
    if mutation == "omit-tail":
        group["block_ids"].pop()
    elif mutation == "duplicate":
        group["block_ids"].append(blocks[0].id)
    elif mutation == "reorder":
        group["block_ids"][1:3] = reversed(group["block_ids"][1:3])
    elif mutation == "invent-heading":
        group["heading_block_id"] = "Invented title"
    else:
        group["heading_block_id"] = blocks[-1].id
    with pytest.raises(ValueError):
        document_from_plan(
            value,
            blocks,
            sha256=hashlib.sha256(SOURCE.encode()).hexdigest(),
            language="zh-CN",
            skill_id="article-format",
        )


def test_document_unicode_offsets_and_sha_cannot_rebind_a_different_source():
    blocks = source_blocks(SOURCE)
    with pytest.raises(ValueError):
        document_from_plan(
            plan(blocks),
            blocks,
            sha256="a" * 64,
            language="en-US",
            skill_id="article-format",
        )


def test_document_storage_keeps_meaningful_leading_and_eof_whitespace_exact():
    original = "\n\n" + SOURCE + "\n\n"
    blocks = source_blocks(original)
    value = plan(blocks)
    value["groups"][0]["heading_block_id"] = None
    result = document_from_plan(
        value,
        blocks,
        sha256=hashlib.sha256(original.encode()).hexdigest(),
        language="zh-CN",
        skill_id="article-format",
    )
    restored = analysis_result_from_document(analysis_result_document(result))
    assert "".join(section.body for section in restored.sections) == original
    restored.validate_document_source(hashlib.sha256(original.encode()).hexdigest())


def test_novel_chapters_keep_exact_unicode_and_missing_eof_newline():
    original = (
        "第一章 取件\n林澈想取父亲的信。\n\n第二章 共同拆"
        "信\n林遥说兄妹需共同写地址。é🙂"
    )
    index = index_drama_source(original)
    units = index["units"]
    assert len(units) == 2
    assert units[0]["start"] == 0
    assert units[-1]["end"] == len(original)
    assert "".join(original[item["start"] : item["end"]] for item in units) == original
    assert index["source_sha256"] == hashlib.sha256(original.encode()).hexdigest()
    assert units[1]["title"] == "第二章 共同拆信"


@pytest.mark.parametrize(
    "labels", [("Chapter 1", "Chapter 2"), ("CHAPTER I", "Chapter II")]
)
def test_english_chapters_index_long_full_text_without_unicode_or_eof_changes(labels):
    original = f"{labels[0]} · é🙂\n" + "Existing source. " * 3500
    original += f"\n{labels[1]} · The letter\n" + "Keep every limit. " * 3300
    index = index_drama_source(original)
    units = index["units"]
    assert len(original) > 100000 and len(units) == 2
    assert index["chapter_unit"] == "chapter"
    assert units[0]["start"] == 0 and units[-1]["end"] == len(original)
    assert "".join(original[unit["start"] : unit["end"]] for unit in units) == original
    assert index["source_sha256"] == hashlib.sha256(original.encode()).hexdigest()


@pytest.mark.parametrize("spans", [[(1, 4)], [(0, 3), (2, 4)], [(0, 5)]])
def test_source_units_reject_missing_preface_overlap_and_out_of_bounds(spans):
    units = tuple(
        ScreenplaySceneSource(f"unit-{index}", start, end)
        for index, (start, end) in enumerate(spans)
    )
    with pytest.raises(AnalysisArtifactError):
        validate_screenplay_source("原文全文", units)


async def test_single_call_rejects_omitted_preface_before_provider_resolution(tmp_path):
    text = "前言与正文"
    path = tmp_path / "source.md"
    path.write_text(text)
    repository, loader, resolver, monitor = (AsyncMock() for _ in range(4))
    repository.get_screenplay_source.return_value = SimpleNamespace(
        character_count=len(text),
        scenes=(ScreenplaySceneSource("scene-1", 1, len(text)),),
    )
    loader.materialize.return_value = SimpleNamespace(screenplay=path)

    async def run(operation, **kwargs):
        return await operation()

    monitor.run.side_effect = run
    executor = ScreenplayAnalysisExecutor(
        repository=repository,
        loader=loader,
        resolver=resolver,
        clock=lambda: NOW,
        max_single_call_characters=120000,
    )
    with pytest.raises(AnalysisArtifactError):
        await executor.execute(
            replace(running_job(), result_contract="screenplay-analysis"), monitor
        )
    resolver.resolve_screenplay.assert_not_awaited()
    monitor.step.assert_not_awaited()
    loader.cleanup.assert_awaited_once()


def test_active_methods_preserve_original_form_contract_and_load_bounded_real_modules():
    catalog = BuiltinAnalysisSkillCatalog()
    assert [item.id for item in catalog.list(AnalysisInputKind.VIDEO)] == [
        "video-review",
        "video-breakdown",
    ]
    assert [item.id for item in catalog.list(AnalysisInputKind.SCREENPLAY)] == [
        "screenplay-analysis",
        "article-format",
        "wechat-format",
        "xhs-format",
    ]
    assert catalog.list(AnalysisInputKind.CONTENT) == ()
    for kind in [AnalysisInputKind.VIDEO, AnalysisInputKind.SCREENPLAY]:
        for view in catalog.list(kind):
            skill = catalog.resolve(view.id, kind)
            assert (
                skill.instructions_sha256
                == hashlib.sha256(skill.instructions.encode()).hexdigest()
            )
            assert len(skill.instructions) <= 30000
            assert "Pinned source:" in skill.instructions
            assert view.default_prompt
    assert catalog.resolve("screenplay-rewrite", AnalysisInputKind.SCREENPLAY) is None
