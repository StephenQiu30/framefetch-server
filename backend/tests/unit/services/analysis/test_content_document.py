import json
from io import BytesIO

import pytest
from app.integrations.analysis_report_docx import PythonDocxAnalysisReportRenderer
from app.services.analysis.content_report import render_content_markdown
from app.services.analysis.rules.content_document import (
    ContentDocumentResult,
    ContentDraft,
    ContentSourceSet,
    content_schema,
)
from app.services.analysis.rules.enums import AnalysisResultKind
from app.services.analysis_execution.content_functions import ContentFunctions
from docx import Document
from pydantic import ValidationError
from tests.unit.workers.analysis.test_content_execution import draft, review, source


def test_short_post_needs_no_title_or_headings():
    value = draft("post")
    value["blocks"] = value["blocks"][:1]
    result = ContentDraft.model_validate(value)
    result.validate_sources(source("post"))
    assert result.title is None


def test_provider_schema_keeps_finite_blocks_without_openapi_only_keywords():
    schema = content_schema(ContentDraft)
    wire = json.dumps(schema)
    assert '"oneOf"' not in wire and '"discriminator"' not in wire
    assert len(schema["properties"]["blocks"]["items"]["anyOf"]) == 4
    assert set(schema["required"]) == set(schema["properties"])
    assert schema["additionalProperties"] is False


@pytest.mark.parametrize("field,value", [("ordered", "false"), ("ordered", 0)])
def test_model_cannot_coerce_list_order(field, value):
    data = draft()
    data["blocks"] = [{"id": "items", "type": "list", field: value, "items": ["正文"]}]
    data["evidence_index"] = []
    with pytest.raises(ValidationError):
        ContentDraft.model_validate(data)


@pytest.mark.parametrize(
    "change", ["unknown_material", "wrong_quote", "unknown_block", "style_as_fact"]
)
def test_citations_must_point_to_exact_original_source(change):
    data = draft()
    citation = data["evidence_index"][0]
    materials = source()
    if change == "unknown_material":
        citation["material_id"] = "outside"
    elif change == "wrong_quote":
        citation["quote"] = "长期防漏"
    elif change == "unknown_block":
        citation["block_id"] = "outside"
    else:
        materials = ContentSourceSet.model_validate(
            {
                **materials.model_dump(),
                "materials": [
                    *materials.model_dump()["materials"],
                    {
                        "id": "style",
                        "title": "作者范文",
                        "text": "长期防漏",
                        "role": "author_style",
                    },
                ],
            }
        )
        citation["material_id"] = "style"
        citation["quote"] = "长期防漏"
    with pytest.raises(ValueError):
        ContentDraft.model_validate(data).validate_sources(materials)


def test_rendered_article_and_docx_exclude_review_and_evidence():
    result = ContentDocumentResult.model_validate(
        {
            **draft(),
            "kind": "content_document",
            "source_set_ref": source().sha256,
            "review_status": "needs_material",
            "review_history": [review(needs_material=True)],
        }
    )
    markdown = render_content_markdown(result)
    document = Document(
        BytesIO(
            PythonDocxAnalysisReportRenderer().render(
                markdown, result_kind=AnalysisResultKind.CONTENT_DOCUMENT
            )
        )
    )
    visible = "\n".join(p.text for p in document.paragraphs)
    for forbidden in (
        "segment-000",
        "notes",
        "收窄到这次观察",
        "编辑摘要",
        "编辑附录",
        "Video Server",
    ):
        assert forbidden not in markdown
        assert forbidden not in visible
    assert "桌面上没有看到水滴" in visible
    assert all(not p.text for p in document.sections[0].header.paragraphs)
    assert all(not p.text for p in document.sections[0].footer.paragraphs)


def test_source_functions_reject_injected_tools_and_other_tasks():
    tools = ContentFunctions(source())
    assert (
        tools.dispatch(
            {
                "name": "read_material",
                "arguments": {"material_id": "notes", "segment_id": "segment-000"},
            }
        )
        == source().materials[0].text
    )
    for name, material in [
        ("execute_script", "notes"),
        ("read_material", "other-task"),
        ("read_material", "../../secret"),
    ]:
        with pytest.raises(ValueError):
            tools.dispatch(
                {
                    "name": name,
                    "arguments": {"material_id": material, "segment_id": "segment-000"},
                }
            )


def test_claimed_pass_requires_a_passing_independent_review():
    with pytest.raises(ValidationError):
        ContentDocumentResult.model_validate(
            {
                **draft(),
                "kind": "content_document",
                "source_set_ref": source().sha256,
                "review_status": "passed",
                "review_history": [review(needs_material=True)],
            }
        )
