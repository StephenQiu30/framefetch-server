from __future__ import annotations

import csv
import hashlib
import json
from io import BytesIO, StringIO
from zipfile import ZipFile

import pytest
from app.services.creation.exports import (
    ExportAsset,
    ExportDocument,
    export_csv,
    export_document,
    sanitize_html,
)
from docx import Document
from PIL import Image


def _image() -> bytes:
    output = BytesIO()
    Image.new("RGB", (64, 32), "red").save(output, format="PNG")
    return output.getvalue()


def test_document_package_opens_with_version_and_unchanged_assets() -> None:
    picture = _image()
    body = "## 原创中文\n\n明确数字 123，保留引文。\n\n![原创图](assets/own.png)"
    source = ExportDocument(
        "原创稿",
        body,
        "revision-2",
        references=[{"source_id": "source-1", "quote": "明确数字 123"}],
        assets=[ExportAsset("own.png", picture, "image/png", "本人原创测试图")],
    )
    artifact = export_document(source, "zip")
    with ZipFile(BytesIO(artifact.data)) as archive:
        assert set(archive.namelist()) == {
            "content.md",
            "content.html",
            "content.docx",
            "manifest.json",
            "assets/own.png",
        }
        manifest = json.loads(archive.read("manifest.json"))
        assert manifest["revision_id"] == "revision-2"
        assert manifest["references"][0]["source_id"] == "source-1"
        for item in manifest["files"]:
            assert (
                hashlib.sha256(archive.read(item["path"])).hexdigest() == item["sha256"]
            )
        assert body in archive.read("content.md").decode()
        html = archive.read("content.html").decode()
        assert '<img src="assets/own.png"' in html
        opened = Document(BytesIO(archive.read("content.docx")))
        assert opened.core_properties.identifier == "revision-2"
        assert "明确数字 123" in "\n".join(
            paragraph.text for paragraph in opened.paragraphs
        )
        assert archive.read("assets/own.png") == picture
        with Image.open(BytesIO(archive.read("assets/own.png"))) as image:
            image.load()
            assert image.size == (64, 32)
    assert source.body == body and source.assets[0].data == picture
    assert artifact.metadata["revision_id"] == "revision-2"


@pytest.mark.parametrize(
    "name", ["../x.png", "/x.png", r"..\x.png", "%2e%2e.png", "a\nx.png"]
)
def test_package_rejects_unsafe_archive_paths(name: str) -> None:
    source = ExportDocument(
        "原创",
        "正文",
        "revision-1",
        assets=[ExportAsset(name, _image(), "image/png", "原创")],
    )
    with pytest.raises(ValueError, match="文件名"):
        export_document(source, "zip")


def test_package_rejects_fake_images_unconfirmed_rights_and_name_collision() -> None:
    for assets, message in [
        ([ExportAsset("x.png", b"not a picture", "image/png", "原创")], "图片"),
        ([ExportAsset("x.webp", _image(), "image/webp", "原创")], "编码"),
        ([ExportAsset("x.png", _image(), "image/png", "unknown")], "权利"),
        ([ExportAsset("x.png", _image(), "image/png", "原创")] * 2, "重复"),
    ]:
        with pytest.raises(ValueError, match=message):
            export_document(ExportDocument("原创", "正文", "r1", assets=assets), "zip")


def test_html_sanitizer_removes_active_content_and_remote_resources() -> None:
    original = (
        '<p onclick="steal()" style="background:url(http://bad)">安全'
        "<script>steal()</script><svg><script>steal()</script></svg>"
        '<a href="javascript:alert(1)">坏链接</a><img src="https://bad/x.png">'
        '<img src="assets/x.png" onerror="steal()"><iframe src="/secret"></iframe>'
        '<a href="https://example.com/guide">来源</a></p>'
    )
    output = sanitize_html(original, allowed_assets=frozenset({"assets/x.png"}))
    for forbidden in (
        "onclick",
        "onerror",
        "<script",
        "<svg",
        "<iframe",
        "javascript:",
        "https://bad",
    ):
        assert forbidden not in output
    assert 'src="assets/x.png"' in output
    assert 'href="https://example.com/guide"' in output
    assert "安全" in output and "坏链接" in output


def test_markdown_html_escapes_raw_markup_and_disallows_remote_images() -> None:
    source = ExportDocument(
        "<script>title</script>",
        "<script>alert(1)</script>\n\n![坏图](https://bad/x.png)",
        "r1",
    )
    output = export_document(source, "html").data.decode()
    assert "<script>" not in output
    assert "https://bad/x.png" not in output
    assert "Content-Security-Policy" in output
    assert "图片未包含" in output


def test_csv_is_readable_unicode_and_formula_safe() -> None:
    result = export_csv(
        [{"镜头": 1, "备注": '=HYPERLINK("bad")'}, {"镜头": 2, "备注": "中文,引文"}],
        ["镜头", "备注"],
    )
    rows = list(csv.DictReader(StringIO(result.data.decode("utf-8-sig"))))
    assert rows[0]["备注"].startswith("'=HYPERLINK")
    assert rows[1] == {"镜头": "2", "备注": "中文,引文"}


def test_zip_never_marks_missing_or_remote_assets_complete() -> None:
    document = ExportDocument("原创", "![缺图](assets/missing.png)", "r1")
    preview = export_document(document, "html")
    assert preview.metadata["resource_status"] == "needs_assets"
    assert preview.metadata["missing_asset_count"] == 1
    with pytest.raises(ValueError, match="缺失"):
        export_document(document, "zip")
