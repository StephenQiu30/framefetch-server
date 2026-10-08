from dataclasses import replace
from datetime import timedelta
from io import BytesIO
from uuid import uuid4

import pytest
from app.integrations.analysis_report_docx import PythonDocxAnalysisReportRenderer
from app.integrations.object_storage import StoredObjectStat
from app.models import ArtifactRow
from app.repositories.analysis.report_repository import (
    ReportObject,
    SqlAlchemyAnalysisReportRepository,
)
from app.repositories.analysis.repository import SqlAlchemyAnalysisRepository
from app.services.analysis.models import AnalysisPublish
from app.services.analysis.rules.enums import AnalysisResultContract
from app.services.quotas import QuotaExceeded, QuotaPolicy
from app.workers.report.message import ReportRequested
from app.workers.report.publisher import ReportPublisher
from docx import Document
from sqlalchemy import func, select
from tests.unit.repositories.analysis.factories import analysis_result
from tests.unit.repositories.analysis.test_publish import NOW, validating_job
from tests.unit.services.analysis.test_editorial_examples import _parse, _read


async def publishing(analysis_db):
    command, job = await validating_job(analysis_db)
    snapshot = await analysis_db.repository.publish_result(
        AnalysisPublish(
            job_id=command.id,
            run_id=command.run_id,
            result=analysis_result(),
            lease_owner="worker-a",
            expected_version=job.version,
            provider="codex",
            model="controlled-model",
            cli_version="controlled",
            now=NOW + timedelta(seconds=3),
        )
    )
    repo = SqlAlchemyAnalysisReportRepository(analysis_db.sessions)
    report = await repo.get_latest_report(command.id)
    return command, snapshot, repo, report


async def test_cancelled_publisher_retains_quota_until_objects_are_deleted(analysis_db):
    command, snapshot, reports, report = await publishing(analysis_db)
    publication = await reports.claim(
        report_id=report.id,
        job_id=command.id,
        run_id=command.run_id,
        expected_version=snapshot.version,
        worker_id="publisher",
        now=NOW + timedelta(seconds=4),
        lease_for=timedelta(minutes=1),
    )
    async with analysis_db.sessions() as session:
        stored = await session.scalar(select(func.sum(ArtifactRow.size_bytes)))
    policy = QuotaPolicy(storage_bytes=stored + QuotaPolicy().report_bytes)
    jobs = SqlAlchemyAnalysisRepository(analysis_db.sessions, quota_policy=policy)
    await jobs.cancel_job(command.id, command.owner_hash, NOW + timedelta(seconds=5))
    second = replace(
        command,
        id=uuid4(),
        run_id=uuid4(),
        outbox_event_id=uuid4(),
        idempotency_key="second",
    )
    with pytest.raises(QuotaExceeded, match="storage_quota_exceeded"):
        await jobs.create_job_and_enqueue(second, now=NOW + timedelta(seconds=6))
    objects = tuple(
        ReportObject(
            format=kind,
            bucket="video-artifacts",
            object_key=f"analyses/{command.id}/runs/1/reports/{report.id}/report.{suffix}",
            content_type=content_type,
            size_bytes=32,
            sha256="a" * 64,
        )
        for kind, suffix, content_type in (
            ("markdown", "md", "text/markdown"),
            ("docx", "docx", "application/octet-stream"),
        )
    )
    await reports.complete(
        publication, "publisher", objects, NOW + timedelta(seconds=7)
    )
    cancelled = await jobs.get_job(command.id)
    assert cancelled.status == "cancelled"
    assert cancelled.current_report_id is None
    assert await reports.get_current_report_file(command.id, "docx") is None
    with pytest.raises(QuotaExceeded, match="storage_quota_exceeded"):
        await jobs.create_job_and_enqueue(second, now=NOW + timedelta(seconds=8))
    deleted = []

    async def delete_object(key):
        deleted.append(key)

    await reports.purge_report_artifacts(NOW + timedelta(seconds=9), delete_object)
    assert set(deleted) == {item.object_key for item in objects}
    assert (
        await jobs.create_job_and_enqueue(second, now=NOW + timedelta(seconds=10))
    ).created


async def test_report_size_rejection_is_terminal_and_cleanup_releases_reservation(
    analysis_db,
):
    command, snapshot, reports, report = await publishing(analysis_db)

    class Storage:
        async def stat(self, key):
            raise AssertionError("oversized report must be rejected before upload")

    publisher = ReportPublisher(
        reports,
        Storage(),
        PythonDocxAnalysisReportRenderer(),
        bucket="video-artifacts",
        worker_id="publisher",
        clock=lambda: NOW + timedelta(seconds=4),
        max_bytes=1024,
    )
    requested = ReportRequested(
        command.id, command.run_id, report.id, report.renderer_version, snapshot.version
    )
    assert await publisher.execute(requested) is True
    job = await analysis_db.repository.get_job(command.id)
    assert job.status == "failed"
    assert job.error_code == "analysis_resource_limit"
    assert job.stage is None
    assert await reports.recover_pending(NOW + timedelta(minutes=10)) == ()
    assert await publisher.execute(requested) is True  # stale redelivery is a no-op
    deleted = []

    async def delete_object(key):
        deleted.append(key)

    result = await reports.purge_report_artifacts(
        NOW + timedelta(minutes=10), delete_object
    )
    assert result.failed == 0
    assert len(deleted) == 2
    assert (await reports.get_latest_report(command.id)).status == "deleted"


async def test_article_publication_exports_clean_copy_and_keeps_review_data(
    analysis_db,
):
    command, job = await validating_job(
        analysis_db, result_contract=AnalysisResultContract.VIDEO_ARTICLE
    )
    article = _parse("video-article", _read("video-article"))
    snapshot = await analysis_db.repository.publish_result(
        AnalysisPublish(
            job_id=command.id,
            run_id=command.run_id,
            result=article,
            lease_owner="worker-a",
            expected_version=job.version,
            provider="codex",
            model="controlled-model",
            cli_version="controlled",
            now=NOW + timedelta(seconds=3),
        )
    )
    reports = SqlAlchemyAnalysisReportRepository(analysis_db.sessions)
    report = await reports.get_latest_report(command.id)
    content_by_key = {}
    stats_by_key = {}

    class Storage:
        async def stat(self, key):
            return stats_by_key.get(key)

        async def upload_bytes(self, key, content, media_type, digest):
            content_by_key[key] = content
            stats_by_key[key] = StoredObjectStat(len(content), digest, media_type)

    publisher = ReportPublisher(
        reports,
        Storage(),
        PythonDocxAnalysisReportRenderer(),
        bucket="video-artifacts",
        worker_id="publisher",
        clock=lambda: NOW + timedelta(seconds=4),
    )
    requested = ReportRequested(
        command.id, command.run_id, report.id, report.renderer_version, snapshot.version
    )
    assert await publisher.execute(requested) is True
    completed = await analysis_db.repository.get_job(command.id)
    assert completed.status == "succeeded"
    assert await analysis_db.repository.get_result(command.id) == article
    assert len(content_by_key) == 3
    for key, content in content_by_key.items():
        if key.endswith(".md"):
            text = content.decode("utf-8")
        elif key.endswith(".html"):
            from html.parser import HTMLParser

            class TextParser(HTMLParser):
                def __init__(self):
                    super().__init__()
                    self.chunks = []

                def handle_data(self, data):
                    self.chunks.append(data)

            parser = TextParser()
            parser.feed(content.decode("utf-8"))
            text = " ".join(parser.chunks).strip()
        else:
            document = Document(BytesIO(content))
            assert document.core_properties.subject == "Article"
            assert not document.sections[0].header.tables
            text = "\n".join(p.text for p in document.paragraphs)
        assert text.rstrip().endswith(article.closing)
        assert "还需要更完整的测试" in text
        for metadata in ("编辑摘要", "编辑附录", "00:07.000", "未使用可靠音频"):
            assert metadata not in text


async def test_native_bundle_is_inserted_after_its_report_and_exposed_in_metadata(
    analysis_db,
):
    from app.models import AnalysisJobRow
    from app.services.analysis.models import AnalysisReportArtifactSnapshot
    from app.services.analysis.rules.result_models import AnalysisMedia
    from app.services.analysis.rules.structured_report import (
        StructuredReportResult,
        StructuredReportSection,
    )

    command, job = await validating_job(
        analysis_db, result_contract=AnalysisResultContract.STRUCTURED_REPORT
    )
    async with analysis_db.sessions() as session, session.begin():
        row = await session.get(AnalysisJobRow, command.id)
        row.skill_id = "video-shots"
    native = AnalysisReportArtifactSnapshot(
        "zip",
        f"analyses/{job.id}/runs/{job.run_id}/native/{'a' * 64}.zip",
        "application/zip",
        100,
        "a" * 64,
    )
    result = StructuredReportResult(
        "zh-CN",
        "拉片",
        "已实测",
        (
            StructuredReportSection(
                "shot-001", "第一镜", "黑色背景上的白色备份说明文字。", (), ()
            ),
        ),
        (),
        AnalysisMedia(2000, "mp4", 100),
    )
    await analysis_db.repository.publish_result(
        AnalysisPublish(
            job_id=command.id,
            run_id=command.run_id,
            result=result,
            lease_owner="worker-a",
            expected_version=job.version,
            provider="codex",
            model="controlled",
            cli_version="controlled",
            now=NOW + timedelta(seconds=3),
            native_markdown="# 原生拉片\n",
            native_artifacts=(native,),
            native_bucket="video-artifacts",
        )
    )
    report = await analysis_db.repository.get_latest_report(command.id)
    assert report.markdown == "# 原生拉片\n"
    assert report.artifacts == (native,)
    snapshot = await analysis_db.repository.get_job(command.id)
    reports = SqlAlchemyAnalysisReportRepository(analysis_db.sessions)
    publication = await reports.claim(
        report_id=report.id,
        job_id=command.id,
        run_id=command.run_id,
        expected_version=snapshot.version,
        worker_id="publisher",
        now=NOW + timedelta(seconds=4),
        lease_for=timedelta(minutes=1),
    )
    assert publication.native_artifact_bytes == native.size_bytes

    class ClaimedRepository:
        async def claim(self, **kwargs):
            return publication

        async def fail(self, *args, **kwargs):
            await reports.fail(*args, **kwargs)

    class Storage:
        async def stat(self, key):
            raise AssertionError("combined native budget must fail before upload")

    class Renderer:
        def render(self, markdown, *, result_kind):
            return b"docx"

    publisher = ReportPublisher(
        ClaimedRepository(),
        Storage(),
        Renderer(),
        bucket="video-artifacts",
        worker_id="publisher",
        clock=lambda: NOW + timedelta(seconds=4),
        max_bytes=native.size_bytes,
    )
    requested = ReportRequested(
        command.id, command.run_id, report.id, report.renderer_version, snapshot.version
    )
    assert await publisher.execute(requested) is True
    failed = await analysis_db.repository.get_job(command.id)
    assert failed.status == "failed"
    assert failed.error_code == "analysis_resource_limit"
